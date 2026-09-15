from pathlib import Path
from datetime import datetime

from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from docx import Document


BASE_DIR = Path(__file__).resolve().parent.parent
REPORT_DIR = BASE_DIR / "reports"

REPORT_DIR.mkdir(exist_ok=True)

HUMAN_VERIFICATION_CHECKS = [
    ("Font Size", "Verify required declarations and numerals meet the applicable minimum size."),
    ("Principal Display Panel", "Confirm required declarations appear on the correct display panel."),
    ("Readability and Contrast", "Confirm the printed information is legible, prominent and sufficiently contrasting."),
    ("Inspection Evidence", "Record any physical inspection, enforcement evidence or supporting documents."),
    ("Deceptive Packaging", "Check that the package appearance does not mislead or exaggerate the declared quantity."),
    ("Wholesale Package", "Confirm whether wholesale-package requirements apply and verify them when applicable."),
    ("Export Package", "Confirm whether export-package requirements apply and verify them when applicable."),
    ("Registration", "Verify applicable manufacturer, packer or importer registration with the relevant authority.")
]


def generate_report_id():
    return datetime.now().strftime("RPT-%Y%m%d-%H%M%S")


def generate_pdf_report(scan_data):
    report_id = generate_report_id()
    file_path = REPORT_DIR / f"{report_id}.pdf"

    pdf = canvas.Canvas(str(file_path), pagesize=A4)

    width, height = A4
    y = height - 50

    pdf.setFont("Helvetica-Bold", 20)
    pdf.drawString(50, y, "Packaged Commodity Compliance Report")

    y -= 35

    pdf.setFont("Helvetica", 10)
    pdf.drawString(50, y, f"Report ID: {report_id}")

    y -= 18

    pdf.drawString(
        50,
        y,
        f"Generated: {datetime.now().strftime('%d-%m-%Y %H:%M:%S')}"
    )

    y -= 35

    product_name = scan_data.get("product_name", "Not provided")
    scan_id = scan_data.get("scan_id", "Not provided")

    pdf.setFont("Helvetica-Bold", 12)
    pdf.drawString(50, y, "Inspection Details")

    y -= 20

    pdf.setFont("Helvetica", 10)

    pdf.drawString(
        50,
        y,
        f"Scan ID: {scan_id}"
    )

    y -= 16

    pdf.drawString(
        50,
        y,
        f"Product: {product_name}"
    )

    y -= 30

    compliance = scan_data.get("compliance", {})

    pdf.setFont("Helvetica-Bold", 12)
    pdf.drawString(50, y, "Compliance Result")

    y -= 20

    pdf.setFont("Helvetica", 10)

    pdf.drawString(
        50,
        y,
        f"Status: {compliance.get('overall_status', 'N/A')}"
    )

    y -= 16

    pdf.drawString(
        50,
        y,
        f"Score: {compliance.get('score', 0)}%"
    )

    y -= 30

    pdf.setFont("Helvetica-Bold", 12)
    pdf.drawString(50, y, "Rule Evaluation")

    y -= 20

    pdf.setFont("Helvetica", 9)

    results = compliance.get("results", [])

    for rule in results:

        status = rule.get("status", "N/A")
        rule_id = rule.get("rule_id", "N/A")
        field = rule.get("field", "N/A")

        line = f"{rule_id} | {field} | {status}"

        pdf.drawString(50, y, line)

        y -= 15

        if y < 60:
            pdf.showPage()
            y = height - 50
            pdf.setFont("Helvetica", 9)

    y -= 15

    pdf.setFont("Helvetica-Bold", 12)
    pdf.drawString(50, y, "Human Verification Instructions")
    y -= 20
    pdf.setFont("Helvetica", 9)

    for check, instruction in HUMAN_VERIFICATION_CHECKS:
        pdf.drawString(50, y, f"{check}: {instruction}")
        y -= 14

        if y < 60:
            pdf.showPage()
            y = height - 50
            pdf.setFont("Helvetica", 9)

    y -= 15

    readability = scan_data.get("readability", {})

    pdf.setFont("Helvetica-Bold", 12)
    pdf.drawString(50, y, "Readability Analysis")

    y -= 20

    pdf.setFont("Helvetica", 10)

    pdf.drawString(
        50,
        y,
        f"Status: {readability.get('status', 'N/A')}"
    )

    y -= 16

    pdf.drawString(
        50,
        y,
        f"OCR Confidence: "
        f"{readability.get('average_confidence', 0)}%"
    )

    y -= 30

    pdf.setFont("Helvetica-Bold", 12)
    pdf.drawString(50, y, "OCR Extracted Text")

    y -= 20

    pdf.setFont("Helvetica", 9)

    extracted_text = scan_data.get(
        "extracted_text",
        "No text detected."
    )

    for line in extracted_text.splitlines():

        if not line.strip():
            continue

        pdf.drawString(
            50,
            y,
            line[:100]
        )

        y -= 13

        if y < 60:
            pdf.showPage()
            y = height - 50
            pdf.setFont("Helvetica", 9)

    pdf.save()

    return str(file_path)


def generate_word_report(scan_data):
    report_id = generate_report_id()
    file_path = REPORT_DIR / f"{report_id}.docx"

    document = Document()

    document.add_heading(
        "Packaged Commodity Compliance Report",
        level=1
    )

    document.add_paragraph(
        f"Report ID: {report_id}"
    )

    document.add_paragraph(
        f"Generated: "
        f"{datetime.now().strftime('%d-%m-%Y %H:%M:%S')}"
    )

    document.add_heading(
        "Inspection Details",
        level=2
    )

    document.add_paragraph(
        f"Scan ID: "
        f"{scan_data.get('scan_id', 'Not provided')}"
    )

    document.add_paragraph(
        f"Product: "
        f"{scan_data.get('product_name', 'Not provided')}"
    )

    compliance = scan_data.get("compliance", {})

    document.add_heading(
        "Compliance Result",
        level=2
    )

    document.add_paragraph(
        f"Status: "
        f"{compliance.get('overall_status', 'N/A')}"
    )

    document.add_paragraph(
        f"Score: "
        f"{compliance.get('score', 0)}%"
    )

    document.add_heading(
        "Rule Evaluation",
        level=2
    )

    results = compliance.get("results", [])

    if results:

        table = document.add_table(
            rows=1,
            cols=3
        )

        table.style = "Table Grid"

        header = table.rows[0].cells

        header[0].text = "Rule"
        header[1].text = "Requirement"
        header[2].text = "Status"

        for rule in results:

            row = table.add_row().cells

            row[0].text = str(
                rule.get("rule_id", "")
            )

            row[1].text = str(
                rule.get("field", "")
            )

            row[2].text = str(
                rule.get("status", "")
            )

    document.add_heading(
        "Human Verification Instructions",
        level=2
    )

    for check, instruction in HUMAN_VERIFICATION_CHECKS:
        document.add_paragraph(
            f"{check}: {instruction}",
            style="List Bullet"
        )

    readability = scan_data.get(
        "readability",
        {}
    )

    document.add_heading(
        "Readability Analysis",
        level=2
    )

    document.add_paragraph(
        f"Status: "
        f"{readability.get('status', 'N/A')}"
    )

    document.add_paragraph(
        f"OCR Confidence: "
        f"{readability.get('average_confidence', 0)}%"
    )

    document.add_heading(
        "OCR Extracted Text",
        level=2
    )

    document.add_paragraph(
        scan_data.get(
            "extracted_text",
            "No text detected."
        )
    )

    document.save(str(file_path))

    return str(file_path)